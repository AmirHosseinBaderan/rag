from pathlib import Path

from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import JSONResponse

from app.domain.documents.document import Document
from app.api.dependencies import get_indexer

DOCUMENTS_DIR = Path(__file__).parents[3] / "documents"
ALLOWED_EXTENSIONS = {".md", ".docx", ".pdf"}

documents_router = APIRouter()


@documents_router.post("")
async def add_document(file: UploadFile = File(...)):
    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="File type not allowed. Use .md, .docx, or .pdf",
        )

    documents_dir = DOCUMENTS_DIR
    documents_dir.mkdir(parents=True, exist_ok=True)
    file_path = documents_dir / file.filename

    with open(file_path, "wb") as f:
        f.write(await file.read())

    content = _extract_text(file_path)

    document = Document(
        id=file_path.stem,
        content=content,
        metadata={"source": file.filename},
    )

    indexer = get_indexer()
    indexer.index(document)

    return JSONResponse(
        status_code=201,
        content={
            "message": "Document indexed successfully",
            "filename": file.filename,
            "document_id": document.id,
        },
    )


@documents_router.get("")
def list_documents():
    documents_dir = DOCUMENTS_DIR
    files = [
        {
            "name": f.name,
            "size": f.stat().st_size,
            "modified": f.stat().st_mtime,
        }
        for f in sorted(documents_dir.iterdir())
        if f.is_file()
    ]
    return {"documents": files}


def _extract_text(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    if ext == ".md":
        return file_path.read_text(encoding="utf-8")
    if ext == ".docx":
        import docx

        doc = docx.Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs)
    if ext == ".pdf":
        from PyPDF2 import PdfReader

        reader = PdfReader(str(file_path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    raise HTTPException(
        status_code=400,
        detail=f"Unsupported file type: {ext}",
    )
